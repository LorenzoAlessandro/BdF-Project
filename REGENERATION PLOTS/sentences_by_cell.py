"""
Dump every test sentence into Excel, one sheet per confusion-matrix cell, with each
model's votes next to it, so the misclassifications can actually be read.

Inputs (both sit in results/cross_eval/FED-v5_x_ECB-v5/):
  cross_eval_predictions.csv.gz   train_institution, model, run, eval_institution,
                                  row_idx, y_true, y_pred
  cross_eval_test_sets.csv.gz     eval_institution, row_idx, text, label_id

The predictions file holds hard labels only (no probabilities), so the "score" is the
share of fits (model x run) that voted for the majority label: 1.00 means every model in
every run made the same call. In the error sheets that is the interesting end, i.e. the
sentences every model gets wrong the same way.

Output: one workbook per (train institution -> eval institution) direction, e.g.
  sentences_FED_model_on_ECB_test.xlsx
with these sheets:
  overview           the 3x3 matrix twice: sentences by majority vote, and pooled fit counts
                     (the second one matches the numbers in the confusion-matrix figures)
  all sentences      every sentence with every column, for filtering/sorting
  dovish as hawkish  ... one sheet per cell, "<true> as <predicted>", sorted by score desc

    python sentences_by_cell.py --preds ".../cross_eval_predictions.csv.gz" --out sentences_v5
    python sentences_by_cell.py --preds ... --models roberta-large-dapt   # one model, votes out of 5 runs
    python sentences_by_cell.py --preds ... --one-file                    # all 4 directions in one workbook
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

LABELS = ["dovish", "hawkish", "neutral"]   # ids 0, 1, 2
IDS = [0, 1, 2]
NAME = dict(zip(IDS, LABELS))
INSTITUTIONS = ["FED", "ECB"]
FONT = "Arial"

WIDTHS = {"row_idx": 8, "sentence": 90, "true label": 11, "predicted (majority)": 19,
          "score": 7, "fits": 6, "dovish votes": 12, "hawkish votes": 13,
          "neutral votes": 13, "tie": 6}


# ----------------------------------------------------------------------------- data
def load(preds_path: Path, tests_path: Path) -> pd.DataFrame:
    preds = pd.read_csv(preds_path)
    need = {"train_institution", "model", "run", "eval_institution", "y_true", "y_pred"}
    if missing := need - set(preds.columns):
        raise SystemExit(f"predictions file is missing columns: {sorted(missing)}")
    if "row_idx" not in preds.columns:
        print("WARNING: predictions have no row_idx column; assuming rows are in test-set "
              "order within each (train, model, run, eval) block")
        preds["row_idx"] = preds.groupby(
            ["train_institution", "model", "run", "eval_institution"]).cumcount()

    tests = pd.read_csv(tests_path)
    if "label_id" not in tests.columns and "label" in tests.columns:
        tests = tests.rename(columns={"label": "label_id"})
    need = {"eval_institution", "row_idx", "text", "label_id"}
    if missing := need - set(tests.columns):
        raise SystemExit(f"test-set file is missing columns: {sorted(missing)}")
    tests = (tests[["eval_institution", "row_idx", "text", "label_id"]]
             .drop_duplicates(["eval_institution", "row_idx"]))

    df = preds.merge(tests, on=["eval_institution", "row_idx"], how="left", validate="m:1")

    # two alignment checks: every prediction must find its sentence, and the label stored
    # with the prediction must equal the label in the archived test set
    if df["text"].isna().any():
        bad = df.loc[df["text"].isna(), ["eval_institution", "row_idx"]].drop_duplicates()
        raise SystemExit(f"{len(bad)} (eval_institution, row_idx) pairs have no sentence in "
                         f"the test-set file, e.g.\n{bad.head().to_string(index=False)}")
    mism = df["y_true"] != df["label_id"]
    if mism.any():
        raise SystemExit(f"{int(mism.sum()):,} of {len(df):,} predictions have y_true != "
                         "label_id after joining on row_idx: the two files don't line up "
                         "(test sets from a different experiment?)")
    return df


# ----------------------------------------------------------------------------- votes
def vote_table(g: pd.DataFrame, models: list) -> pd.DataFrame:
    """One row per sentence: votes per class pooled over the selected model x run fits,
    the majority label and its share (= score), plus a per-model breakdown."""
    votes = (g.groupby(["row_idx", "y_pred"]).size().unstack(fill_value=0)
              .reindex(columns=IDS, fill_value=0))
    n_fits = g.groupby("row_idx").size().reindex(votes.index)
    top = votes.max(axis=1)
    majority = votes.idxmax(axis=1)                  # ties -> first class in IDS order
    tie = votes.eq(top, axis=0).sum(axis=1) > 1
    first = g.drop_duplicates("row_idx").set_index("row_idx").reindex(votes.index)

    out = pd.DataFrame({
        "row_idx": votes.index,
        "sentence": first["text"].values,
        "true label": first["y_true"].map(NAME).values,
        "predicted (majority)": majority.map(NAME).values,
        "score": (top / n_fits).round(3).values,
        "fits": n_fits.values,
        "dovish votes": votes[0].values,
        "hawkish votes": votes[1].values,
        "neutral votes": votes[2].values,
        "tie": tie.values,
    })

    per_model = (g.groupby(["row_idx", "model", "y_pred"]).size().unstack(fill_value=0)
                  .reindex(columns=IDS, fill_value=0))
    present = set(per_model.index.get_level_values("model"))
    for m in models:
        if m not in present:
            continue
        sub = per_model.xs(m, level="model").reindex(votes.index, fill_value=0)
        out[m] = [", ".join(f"{NAME[c]} {int(n)}" for c, n in zip(IDS, row) if n)
                  for row in sub.values]
    return out


def matrices(table: pd.DataFrame, g: pd.DataFrame):
    by_majority = (pd.crosstab(table["true label"], table["predicted (majority)"])
                     .reindex(index=LABELS, columns=LABELS, fill_value=0))
    by_fits = (pd.crosstab(g["y_true"].map(NAME), g["y_pred"].map(NAME))
                 .reindex(index=LABELS, columns=LABELS, fill_value=0))
    for m in (by_majority, by_fits):
        m.index.name, m.columns.name = "true \\ predicted", None
    return by_majority, by_fits


# ----------------------------------------------------------------------------- excel
def style_sheet(ws, columns, wrap_col="sentence"):
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True)
    for i, col in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 24)
    wrap_idx = columns.index(wrap_col) + 1 if wrap_col in columns else None
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(cell.column == wrap_idx))
    ws.freeze_panes = "C2"                          # header row + row_idx/sentence stay put
    ws.auto_filter.ref = ws.dimensions


def write_overview(writer, sheet, tr, ev, models, n_runs, by_majority, by_fits):
    header = pd.DataFrame({"": [
        f"{tr}-trained models evaluated on the {ev} test set",
        f"models: {', '.join(models)}  |  runs per model: {n_runs}  |  "
        f"fits per sentence: {len(models) * n_runs}",
        "score = share of fits voting for the majority label (no probabilities are stored, "
        "only hard predictions)",
        "cell sheets are named '<true> as <predicted>' and hold the sentences whose majority "
        "vote lands in that cell",
    ]})
    header.to_excel(writer, sheet_name=sheet, index=False, header=False)
    r = len(header) + 2
    pd.DataFrame({"": ["sentences by majority vote (each sentence counted once)"]}).to_excel(
        writer, sheet_name=sheet, index=False, header=False, startrow=r)
    by_majority.to_excel(writer, sheet_name=sheet, startrow=r + 1)
    r += len(by_majority) + 4
    pd.DataFrame({"": ["pooled fit counts (every model x run counted; matches the figures)"]}
                 ).to_excel(writer, sheet_name=sheet, index=False, header=False, startrow=r)
    by_fits.to_excel(writer, sheet_name=sheet, startrow=r + 1)

    ws = writer.sheets[sheet]
    ws.column_dimensions["A"].width = 18
    for col in "BCD":
        ws.column_dimensions[col].width = 11
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=isinstance(cell.value, str)
                             and cell.value in LABELS + ["true \\ predicted"])


def write_direction(writer, tr, ev, g, models, prefix=""):
    table = vote_table(g, models)
    by_majority, by_fits = matrices(table, g)
    n_runs = int(g["run"].nunique())
    write_overview(writer, f"{prefix}overview", tr, ev, models, n_runs, by_majority, by_fits)

    table = table.sort_values(["true label", "predicted (majority)", "score"],
                              ascending=[True, True, False])
    cols = list(table.columns)
    table.to_excel(writer, sheet_name=f"{prefix}all sentences", index=False)
    style_sheet(writer.sheets[f"{prefix}all sentences"], cols)

    for t in LABELS:
        for p in LABELS:
            cell = (table[(table["true label"] == t) & (table["predicted (majority)"] == p)]
                    .sort_values(["score", "row_idx"], ascending=[False, True]))
            name = f"{prefix}{t} as {p}"
            cell.to_excel(writer, sheet_name=name, index=False)
            style_sheet(writer.sheets[name], cols)
    return by_majority, by_fits


# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--preds", required=True, help="path to cross_eval_predictions.csv.gz")
    ap.add_argument("--test-sets", default=None,
                    help="path to cross_eval_test_sets.csv.gz "
                         "(default: cross_eval_test_sets.csv.gz next to --preds)")
    ap.add_argument("--out", default="sentences_v5", help="output directory")
    ap.add_argument("--models", nargs="*", default=None,
                    help="subset of model names (default: all models in the file)")
    ap.add_argument("--one-file", action="store_true",
                    help="put all four directions in a single workbook "
                         "(sheet names prefixed FED>ECB etc.)")
    args = ap.parse_args()

    preds_path = Path(args.preds)
    tests_path = (Path(args.test_sets) if args.test_sets
                  else preds_path.parent / "cross_eval_test_sets.csv.gz")
    if not tests_path.exists():
        raise SystemExit(f"test-set file not found: {tests_path}\n"
                         "pass it explicitly with --test-sets")

    df = load(preds_path, tests_path)
    models = args.models or sorted(df["model"].unique())
    if unknown := set(models) - set(df["model"].unique()):
        raise SystemExit(f"models not in the predictions file: {sorted(unknown)}")
    df = df[df["model"].isin(models)]
    print(f"{len(df):,} predictions | models: {', '.join(models)} | "
          f"runs {sorted(int(r) for r in df['run'].unique())}")

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    tag = f"{models[0]}_" if len(models) == 1 else ""
    directions = [(tr, ev) for tr in INSTITUTIONS for ev in INSTITUTIONS]

    if args.one_file:
        path = out / f"sentences_{tag}all_directions.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for tr, ev in directions:
                g = df[(df.train_institution == tr) & (df.eval_institution == ev)]
                if g.empty:
                    continue
                bm, _ = write_direction(writer, tr, ev, g, models, prefix=f"{tr}>{ev} ")
                print(f"  {tr} -> {ev}: {int(bm.values.sum())} sentences")
        print(f"wrote {path}")
        return

    for tr, ev in directions:
        g = df[(df.train_institution == tr) & (df.eval_institution == ev)]
        if g.empty:
            print(f"  {tr} -> {ev}: no predictions, skipped")
            continue
        path = out / f"sentences_{tag}{tr}_model_on_{ev}_test.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            bm, _ = write_direction(writer, tr, ev, g, models)
        off = int(bm.values.sum() - bm.values.diagonal().sum())
        print(f"  {tr} -> {ev}: {int(bm.values.sum())} sentences, {off} off-diagonal by "
              f"majority vote -> {path}")


if __name__ == "__main__":
    main()
